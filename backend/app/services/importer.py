"""Read CSV or Excel uploads into a list of dicts keyed by lower-case header."""

import csv
import io


class ImportError_(ValueError):
    pass


def _norm(header) -> str:
    return str(header or "").strip().lower().replace(" ", "_")


def read_rows(filename: str, content: bytes) -> list[dict]:
    name = filename.lower()
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        return [{_norm(k): (v.strip() if isinstance(v, str) else v) for k, v in row.items()} for row in reader]
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            headers = [_norm(h) for h in next(rows)]
        except StopIteration:
            return []
        out = []
        for values in rows:
            if values is None or all(v in (None, "") for v in values):
                continue
            out.append({h: v for h, v in zip(headers, values, strict=False) if h})
        return out
    raise ImportError_("Upload a .csv or .xlsx file")
